#!/usr/bin/env python3
"""Validate PRD Forge docx output.

Usage:
    python scripts/validate_prd.py <path-to-docx> [--strict]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator, Sequence

try:
    import yaml
    from docx import Document
    from docx.document import Document as DocumentType
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph
except ImportError as exc:  # pragma: no cover - environment guard
    print(f"python-docx and PyYAML are required: {exc}", file=sys.stderr)
    sys.exit(2)


IssueLevel = str


@dataclass
class Issue:
    check_id: str
    level: IssueLevel
    location: str
    matched: str
    suggestion: str

    def to_dict(self) -> dict[str, str]:
        return {
            "check_id": self.check_id,
            "level": self.level,
            "location": self.location,
            "matched": self.matched,
            "suggestion": self.suggestion,
        }


@dataclass
class Block:
    kind: str
    obj: Paragraph | Table
    index: int
    section: str = ""
    heading_text: str = ""
    heading_level: int | None = None


DENY_PATTERNS: Sequence[tuple[str, re.Pattern[str], str]] = (
    ("tool_name", re.compile(r"\b(Codex|ChatGPT|Claude|GPT|Opus|Gemini)\b"), "删除生成工具或模型名称。"),
    ("source_trace", re.compile(r"文档由 AI|由模型|由智能体|由工具生成|本文由 AI"), "删除生成来源痕迹。"),
    ("local_path", re.compile(r"/Users/|file://|\.od/projects/|C:\\"), "删除本地路径或工程目录，只保留文件名、版本号或线上链接。"),
    ("old_structure", re.compile(r"基本信息表|字段｜内容|一、入口：|▌"), "删除旧结构或旧排版残留。"),
)

INTERNAL_ID_RE = re.compile(r"\b[MS]\d{1,2}\b|screen\d+|page\d+", re.I)
TITLE_PREFIX_RE = re.compile(r"^\d+(?:\.\d+)*\.?\s+\S")
HEADING_RE = re.compile(r"^(Heading|标题)\s*([1-5])$")
SECTION_RE = re.compile(r"^(\d+)(?:\.|\s)")
RISK_KEYWORDS = ("资金", "奖励", "风控", "结算", "开奖", "赔率", "抽奖")
LITERAL_PATTERNS = ("<br>", "<br/>", r"\n")
CELL_BREAK_RE = re.compile(r"<(?:\w+:)?br(?:\s+[^>]*)?/?>")
TITLE_END_PUNCTUATION = ("。", "，", "、", "；", "！", "？")


def iter_block_items(parent: DocumentType | _Cell) -> Iterator[Paragraph | Table]:
    if isinstance(parent, DocumentType):
        parent_elm = parent.element.body
        parent_obj = parent
    else:
        parent_elm = parent._tc
        parent_obj = parent

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent_obj)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent_obj)


def cell_image_count(cell: _Cell) -> int:
    # python-docx does not expose images per cell. The XML is stable enough for
    # this validation use case and catches DrawingML pictures inserted in runs.
    return cell._tc.xml.count("<pic:pic")


def table_text(table: Table) -> str:
    parts: list[str] = []
    for row in table.rows:
        for cell in row.cells:
            parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def heading_level_from_style(style_name: str) -> int | None:
    match = HEADING_RE.match(style_name)
    return int(match.group(2)) if match else None


def is_numbered_title(text: str) -> bool:
    stripped = text.strip()
    return (
        bool(TITLE_PREFIX_RE.match(stripped))
        and len(stripped) <= 50
        and not stripped.rstrip().endswith(TITLE_END_PUNCTUATION)
    )


def paragraph_texts(doc: DocumentType) -> list[tuple[str, str, str]]:
    texts: list[tuple[str, str, str]] = []
    p_index = 0
    t_index = 0
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            p_index += 1
            style_name = block.style.name if block.style is not None else ""
            context = "heading" if heading_level_from_style(style_name) is not None else "paragraph"
            texts.append((f"paragraph {p_index}", block.text, context))
        else:
            t_index += 1
            for r_idx, row in enumerate(block.rows, start=1):
                for c_idx, cell in enumerate(row.cells, start=1):
                    for p_idx, paragraph in enumerate(cell.paragraphs, start=1):
                        loc = f"table {t_index} r{r_idx}c{c_idx} p{p_idx}"
                        texts.append((loc, paragraph.text, "table_cell"))
    return texts


def build_blocks(doc: DocumentType) -> list[Block]:
    blocks: list[Block] = []
    current_section = ""
    current_heading = ""
    current_level: int | None = None
    for index, item in enumerate(iter_block_items(doc), start=1):
        if isinstance(item, Paragraph):
            style = item.style.name if item.style is not None else ""
            text = item.text.strip()
            level = heading_level_from_style(style)
            if level and text:
                current_heading = text
                current_level = level
                section_match = SECTION_RE.match(text)
                if current_level == 1 and section_match:
                    current_section = section_match.group(1)
            blocks.append(Block("paragraph", item, index, current_section, current_heading, current_level))
        else:
            blocks.append(Block("table", item, index, current_section, current_heading, current_level))
    return blocks


def add_issue(
    issues: list[Issue],
    check_id: str,
    level: IssueLevel,
    location: str,
    matched: str,
    suggestion: str,
) -> None:
    issues.append(Issue(check_id, level, location, matched, suggestion))


def check_c1(doc: DocumentType, issues: list[Issue]) -> None:
    for location, text, context in paragraph_texts(doc):
        if not text:
            continue
        for name, pattern, suggestion in DENY_PATTERNS:
            for match in pattern.finditer(text):
                add_issue(issues, "C1", "error", location, match.group(0), suggestion)
        if context in ("heading", "table_cell") or not (context == "paragraph" and len(text) > 30):
            for match in INTERNAL_ID_RE.finditer(text):
                add_issue(issues, "C1", "error", location, match.group(0), "删除内部编号，改为中文业务名。")


def check_c2(doc: DocumentType, issues: list[Issue]) -> None:
    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if not is_numbered_title(text):
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if heading_level_from_style(style_name) is None:
            add_issue(
                issues,
                "C2",
                "error",
                f"paragraph {idx}",
                f"{text} (style={style_name or 'None'})",
                "编号正文标题必须使用 Word Heading 1-5，不能用 Normal 加粗伪装。",
            )


def check_c3(doc: DocumentType, issues: list[Issue]) -> None:
    for t_idx, table in enumerate(doc.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            for c_idx, cell in enumerate(row.cells, start=1):
                if CELL_BREAK_RE.search(cell._tc.xml):
                    add_issue(
                        issues,
                        "C3",
                        "error",
                        f"table {t_idx} r{r_idx}c{c_idx}",
                        "w:br",
                        "单元格内含真实换行符，请拆为多个段落。",
                    )
                for p_idx, paragraph in enumerate(cell.paragraphs, start=1):
                    text = paragraph.text
                    for literal in LITERAL_PATTERNS:
                        if literal in text:
                            add_issue(
                                issues,
                                "C3",
                                "error",
                                f"table {t_idx} r{r_idx}c{c_idx} p{p_idx}",
                                literal,
                                "单元格内请拆成多个段落，不要写入换行字面量或 HTML 换行标签。",
                            )


def header_index(table: Table, header_name: str) -> int | None:
    if not table.rows:
        return None
    headers = [cell.text.strip() for cell in table.rows[0].cells]
    for idx, value in enumerate(headers):
        if value == header_name:
            return idx
    return None


def is_feature_table(table: Table) -> bool:
    if not table.rows:
        return False
    headers = {cell.text.strip() for cell in table.rows[0].cells}
    return {"模块", "展示", "需求描述"}.issubset(headers)


def feature_tables(blocks: list[Block]) -> Iterator[tuple[int, str, Table]]:
    table_no = 0
    for block in blocks:
        if block.kind != "table":
            continue
        table_no += 1
        table = block.obj
        if isinstance(table, Table) and block.section == "4" and is_feature_table(table):
            yield table_no, block.heading_text, table


def check_c4(doc: DocumentType, blocks: list[Block], issues: list[Issue]) -> None:
    image_count = len(doc.inline_shapes)
    if image_count == 0:
        add_issue(issues, "C4", "error", "document", "0 images", "PRD 至少应嵌入真实截图，或在展示列标注待设计补图。")

    for table_no, heading, table in feature_tables(blocks):
        display_idx = header_index(table, "展示")
        module_idx = header_index(table, "模块")
        if display_idx is None:
            continue
        for r_idx, row in enumerate(table.rows[1:], start=2):
            display_cell = row.cells[display_idx]
            module_name = row.cells[module_idx].text.strip() if module_idx is not None else heading
            display_text = display_cell.text.strip()
            has_image = cell_image_count(display_cell) > 0
            has_fallback = "待设计补图" in display_text
            if not has_image and not has_fallback:
                add_issue(
                    issues,
                    "C4",
                    "error",
                    f"table {table_no} r{r_idx}c{display_idx + 1}",
                    module_name or heading or "展示列为空",
                    "功能模块展示列必须包含真实图片，或明确写待设计补图。",
                )


def section_text(blocks: list[Block], section: str) -> str:
    parts: list[str] = []
    for block in blocks:
        if block.section != section:
            continue
        if isinstance(block.obj, Paragraph):
            parts.append(block.obj.text)
        elif isinstance(block.obj, Table):
            parts.append(table_text(block.obj))
    return "\n".join(parts)


def admin_config_rows(blocks: list[Block]) -> int:
    count = 0
    for block in blocks:
        if block.section != "5" or not isinstance(block.obj, Table):
            continue
        table = block.obj
        item_idx = header_index(table, "配置项")
        if item_idx is None:
            continue
        for row in table.rows[1:]:
            value = row.cells[item_idx].text.strip()
            if value and value != "/":
                count += 1
    return count


def check_c5(blocks: list[Block], issues: list[Issue], strict: bool) -> None:
    # TODO(P2): Once 管理端可配置 carries field names, replace count checks
    # with unordered field-name set matching between feature details and admin rows.
    feature_text = section_text(blocks, "4")
    mentions = feature_text.count("管理端可配置")
    rows = admin_config_rows(blocks)
    if mentions > rows:
        add_issue(
            issues,
            "C5",
            "error" if strict else "warn",
            "section 4 -> section 5",
            f"管理端可配置={mentions}, 配置项行={rows}",
            "功能详情中的管理端可配置字段必须在管理端需求中有对应配置项。",
        )
    if rows > mentions:
        add_issue(
            issues,
            "C5",
            "error" if strict else "warn",
            "section 5 -> section 4",
            f"配置项行={rows}, 管理端可配置={mentions}",
            "管理端配置项数大于前端引用数，可能存在过度扩展，请核对。",
        )


def collect_heading_sections(blocks: list[Block]) -> list[tuple[Block, str]]:
    sections: list[tuple[Block, str]] = []
    for idx, block in enumerate(blocks):
        if not isinstance(block.obj, Paragraph):
            continue
        style_name = block.obj.style.name if block.obj.style is not None else ""
        level = heading_level_from_style(style_name)
        if not level:
            continue
        text = block.obj.text.strip()
        content: list[str] = []
        for next_block in blocks[idx + 1 :]:
            if isinstance(next_block.obj, Paragraph):
                next_style = next_block.obj.style.name if next_block.obj.style is not None else ""
                next_level = heading_level_from_style(next_style)
                if next_level and next_level <= level:
                    break
                content.append(next_block.obj.text)
            elif isinstance(next_block.obj, Table):
                content.append(table_text(next_block.obj))
        sections.append((block, "\n".join(content)))
    return sections


def check_c6(blocks: list[Block], issues: list[Issue], strict: bool) -> None:
    if not strict:
        return
    for heading, content in collect_heading_sections(blocks):
        title = heading.obj.text.strip() if isinstance(heading.obj, Paragraph) else ""
        if heading.section != "4":
            continue
        if heading.heading_level not in (2, 3):
            continue
        if not any(keyword in title for keyword in RISK_KEYWORDS):
            continue
        missing = [label for label in ("AC", "异常路径", "owner") if label not in content]
        if missing:
            add_issue(
                issues,
                "C6",
                "error",
                title,
                "缺少：" + "、".join(missing),
                "正式评审/归档模式下，高风险模块必须包含 AC、异常路径和 owner。",
            )


def confirmation_rows(blocks: list[Block]) -> list[str]:
    rows: list[str] = []
    for block in blocks:
        if not isinstance(block.obj, Table):
            continue
        table = block.obj
        if not table.rows:
            continue
        headers = {cell.text.strip() for cell in table.rows[0].cells}
        if {"状态", "确认事项"}.issubset(headers):
            for row in table.rows[1:]:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
    return rows


def fallback_modules(blocks: list[Block]) -> list[tuple[str, str]]:
    modules: list[tuple[str, str]] = []
    for table_no, heading, table in feature_tables(blocks):
        display_idx = header_index(table, "展示")
        module_idx = header_index(table, "模块")
        if display_idx is None:
            continue
        for r_idx, row in enumerate(table.rows[1:], start=2):
            if "待设计补图" not in row.cells[display_idx].text:
                continue
            module_name = row.cells[module_idx].text.strip() if module_idx is not None else heading
            modules.append((module_name or heading or "未命名模块", f"table {table_no} r{r_idx}"))
    return modules


def chapter4_module_headings(blocks: list[Block]) -> list[str]:
    modules: list[str] = []
    for block in blocks:
        if block.kind != "paragraph" or block.section != "4":
            continue
        if block.heading_level != 2:
            continue
        if not isinstance(block.obj, Paragraph):
            continue
        text = block.obj.text.strip()
        if text:
            modules.append(text)
    return modules


def feature_display_fallbacks(blocks: list[Block]) -> dict[str, bool]:
    status: dict[str, bool] = {}
    for _table_no, heading, table in feature_tables(blocks):
        display_idx = header_index(table, "展示")
        module_idx = header_index(table, "模块")
        if display_idx is None:
            continue
        for row in table.rows[1:]:
            module_name = row.cells[module_idx].text.strip() if module_idx is not None else heading
            key = module_name or heading
            if key:
                status[key] = "待设计补图" in row.cells[display_idx].text
    return status


def row_has_reason_and_owner(row_text: str) -> bool:
    has_reason = "原因" in row_text or "无法" in row_text or "缺失" in row_text
    has_owner = "负责人" in row_text or "owner" in row_text.lower() or "@" in row_text
    return has_reason and has_owner


def check_c7(blocks: list[Block], issues: list[Issue], strict: bool) -> None:
    rows = confirmation_rows(blocks)
    for module_name, location in fallback_modules(blocks):
        matched_rows = [row for row in rows if "待设计补图" in row and module_name in row]
        if not matched_rows:
            add_issue(
                issues,
                "C7",
                "error" if strict else "warn",
                location,
                module_name,
                "展示列待设计补图必须在确认事项表登记对应模块、缺失原因和负责人；建议在确认事项表的模块列写明完整中文模块名以便交叉定位。",
            )
            continue
        if not any(row_has_reason_and_owner(row) for row in matched_rows):
            add_issue(
                issues,
                "C7",
                "error" if strict else "warn",
                location,
                " | ".join(matched_rows),
                "确认事项登记需同时包含缺失原因和负责人；建议在确认事项表的模块列写明完整中文模块名以便交叉定位。",
            )


def load_manifest(manifest_path: Path, issues: list[Issue], strict: bool) -> dict | None:
    try:
        data = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
    except Exception as exc:
        add_issue(issues, "C8", "error", str(manifest_path), type(exc).__name__, "无法读取 manifest YAML，请检查文件格式。")
        return None
    if not isinstance(data, dict):
        add_issue(issues, "C8", "error", str(manifest_path), "non-object manifest", "manifest 顶层必须是对象。")
        return None
    if not isinstance(data.get("modules"), list):
        add_issue(issues, "C8", "error", "manifest.modules", str(data.get("modules")), "manifest.modules 必须是数组。")
        return None
    return data


def check_c8(blocks: list[Block], issues: list[Issue], strict: bool, manifest_path: Path | None) -> None:
    if manifest_path is None:
        return
    manifest = load_manifest(manifest_path, issues, strict)
    if manifest is None:
        return

    prd_modules = chapter4_module_headings(blocks)
    prd_set = set(prd_modules)
    manifest_modules: list[dict] = [item for item in manifest.get("modules", []) if isinstance(item, dict)]
    manifest_names = [str(item.get("module_name", "")).strip() for item in manifest_modules]
    manifest_set = set(name for name in manifest_names if name)

    missing_in_manifest = sorted(prd_set - manifest_set)
    extra_in_manifest = sorted(manifest_set - prd_set)
    if missing_in_manifest or extra_in_manifest:
        matched = f"PRD缺manifest={missing_in_manifest}; manifest多余={extra_in_manifest}"
        add_issue(
            issues,
            "C8",
            "error",
            "manifest.modules",
            matched,
            "manifest module_name 集合必须与 PRD 第 4 章二级模块标题严格相等。",
        )

    display_fallbacks = feature_display_fallbacks(blocks)
    manifest_missing = {str(item.get("module_name", "")).strip() for item in manifest_modules if item.get("source") == "missing"}
    prd_fallbacks = {name for name, has_fallback in display_fallbacks.items() if has_fallback}
    fallback_mismatch = sorted(manifest_missing.symmetric_difference(prd_fallbacks))
    if fallback_mismatch:
        add_issue(
            issues,
            "C8",
            "error",
            "manifest.modules.source",
            str(fallback_mismatch),
            "source == missing 的模块必须与 PRD 展示列待设计补图双向一致。",
        )

    soft_level = "error" if strict else "warn"
    for idx, item in enumerate(manifest_modules, start=1):
        name = str(item.get("module_name", "")).strip() or f"modules[{idx}]"
        source = item.get("source")
        if source == "missing":
            missing_reason = str(item.get("missing_reason", "")).strip()
            owner = str(item.get("owner", "")).strip()
            if not missing_reason or not owner:
                add_issue(
                    issues,
                    "C8",
                    soft_level,
                    f"manifest.modules[{idx}] {name}",
                    f"missing_reason={missing_reason!r}, owner={owner!r}",
                    "source == missing 时 missing_reason 和 owner 必填。",
                )
        else:
            try:
                dpr = float(item.get("device_pixel_ratio", 0))
            except (TypeError, ValueError):
                dpr = 0.0
            if dpr < 2.0:
                add_issue(
                    issues,
                    "C8",
                    soft_level,
                    f"manifest.modules[{idx}] {name}",
                    f"device_pixel_ratio={item.get('device_pixel_ratio')!r}",
                    "非 missing 模块的 device_pixel_ratio 必须 >= 2.0。",
                )


def validate(path: Path, strict: bool, manifest_path: Path | None = None) -> tuple[list[Issue], dict[str, int | str | bool]]:
    doc = Document(path)
    blocks = build_blocks(doc)
    issues: list[Issue] = []

    check_c1(doc, issues)
    check_c2(doc, issues)
    check_c3(doc, issues)
    check_c4(doc, blocks, issues)
    check_c5(blocks, issues, strict)
    check_c6(blocks, issues, strict)
    check_c7(blocks, issues, strict)
    check_c8(blocks, issues, strict, manifest_path)

    meta: dict[str, int | str | bool] = {
        "path": str(path),
        "strict": strict,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "issues": len(issues),
        "errors": sum(1 for issue in issues if issue.level == "error"),
        "warnings": sum(1 for issue in issues if issue.level == "warn"),
        "manifest": str(manifest_path) if manifest_path else "",
    }
    return issues, meta


def exit_code(issues: Sequence[Issue]) -> int:
    if any(issue.level == "error" for issue in issues):
        return 2
    if any(issue.level == "warn" for issue in issues):
        return 1
    return 0


def print_report(issues: Sequence[Issue], meta: dict[str, int | str | bool]) -> None:
    status = "pass" if not issues else "fail"
    payload = {
        "status": status,
        "meta": meta,
        "issues": [issue.to_dict() for issue in issues],
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    print()
    print("Summary")
    print(f"- file: {meta['path']}")
    print(f"- strict: {meta['strict']}")
    print(f"- tables/images: {meta['tables']}/{meta['images']}")
    print(f"- errors/warnings: {meta['errors']}/{meta['warnings']}")
    if issues:
        for issue in issues:
            print(f"- [{issue.level}] {issue.check_id} {issue.location}: {issue.matched}")
    else:
        print("- all checks passed")


def parse_args(argv: Sequence[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate PRD Forge docx output.")
    parser.add_argument("docx", type=Path, help="Path to .docx file")
    parser.add_argument("--strict", action="store_true", help="Use formal review/archive validation")
    parser.add_argument("--manifest", type=Path, help="Optional screenshot manifest YAML path")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    if not args.docx.exists():
        print(f"File not found: {args.docx}", file=sys.stderr)
        return 2
    issues, meta = validate(args.docx, args.strict, args.manifest)
    print_report(issues, meta)
    return exit_code(issues)


if __name__ == "__main__":
    sys.exit(main())
